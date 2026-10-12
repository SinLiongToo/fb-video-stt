#!/usr/bin/env python3
"""
FB／YouTube 影片語音轉文字：本機伺服器。

瀏覽器沒辦法直接抓 Facebook／YouTube 影片（CORS + 需要解析頁面），所以由這支腳本在本機：
  1. 用 yt-dlp 下載 FB 影片／Reel、YouTube 影片／Shorts 的音軌（audio only，不需要 ffmpeg 合併）
  2. 用 faster-whisper 在本機 CPU 辨識成文字
音檔不會送到任何第三方服務。

用法：
    python server.py [port]
預設 port 8792，啟動後開 http://localhost:8792/
"""
import os
import re
import sys
import tempfile
import threading
import time
import uuid

from flask import Flask, jsonify, request, send_from_directory

if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

HERE = os.path.dirname(os.path.abspath(__file__))
WORK_DIR = os.path.join(tempfile.gettempdir(), 'fb_stt')
os.makedirs(WORK_DIR, exist_ok=True)

# 只接受 Facebook／YouTube 網址，避免變成任意網址下載器
URL_RE = re.compile(
    r'^https?://([a-z0-9-]+\.)?(facebook\.com|fb\.watch|fb\.com|youtube\.com|youtu\.be)/', re.I)
MODELS = {'tiny', 'base', 'small', 'medium', 'large-v3-turbo', 'large-v3'}

# 台語（實驗性）：MediaTek-Research/Breeze-ASR-26 的 CTranslate2 int8 版，給 faster-whisper 用。
# 注意：Breeze 聽台語、但輸出「華語漢字」（官方 model card 明講），不是台語正字。
# 原版 fp32 約 6GB，在 7.7GB RAM 的電腦上直接跑會記憶體不足；int8 版約 1.6GB。
# 下載方式與選這個模型的實測過程見 README「台語模型」。放在 OneDrive 外，避免同步 1.6GB。
TAIGI_MODEL_DIR = os.environ.get(
    'FB_STT_TAIGI_MODEL',
    os.path.join(os.path.expanduser('~'), '.cache', 'fb_stt', 'breeze-asr-26-ct2-int8'))

# GitHub Pages 版的網頁（https://sinliongtoo.github.io/fb-video-stt/）會來呼叫這台本機伺服器。
# CORS 只開給這一個來源；別的網站的 JS 不能操作這台伺服器。
ALLOWED_ORIGINS = {'https://sinliongtoo.github.io'}

app = Flask(__name__)


@app.after_request
def cors(resp):
    origin = request.headers.get('Origin')
    if origin in ALLOWED_ORIGINS:
        resp.headers['Access-Control-Allow-Origin'] = origin
        resp.headers['Access-Control-Allow-Methods'] = 'GET, POST, OPTIONS'
        resp.headers['Access-Control-Allow-Headers'] = 'Content-Type'
        # Chrome 的 Private Network Access：公開網站連 localhost 需要這個
        resp.headers['Access-Control-Allow-Private-Network'] = 'true'
        resp.headers['Vary'] = 'Origin'
    return resp


@app.route('/api/<path:_>', methods=['OPTIONS'])
def preflight(_):
    return '', 204


@app.get('/api/ping')
def ping():
    return jsonify(ok=True)


jobs = {}           # job_id -> dict(status, progress, message, result, error)
_models = {}         # model name -> WhisperModel（只留最近用的一個，RAM 不夠同時放兩個大模型）
_model_lock = threading.Lock()


def get_model(name):
    with _model_lock:
        if name not in _models:
            from faster_whisper import WhisperModel
            _models.clear()
            _models[name] = WhisperModel(name, device='cpu', compute_type='int8')
        return _models[name]


def download_audio(url, job):
    import yt_dlp
    out_tmpl = os.path.join(WORK_DIR, f'{job["id"]}.%(ext)s')

    def hook(d):
        if d['status'] == 'downloading':
            total = d.get('total_bytes') or d.get('total_bytes_estimate') or 0
            if total:
                job['progress'] = int(d.get('downloaded_bytes', 0) * 100 / total)

    opts = {
        # 優先只抓音軌；沒有獨立音軌時退回含音訊的單一 mp4（不需 ffmpeg 合併）
        'format': 'bestaudio/best[acodec!=none]/best',
        'outtmpl': out_tmpl,
        'noplaylist': True,
        'quiet': True,
        'no_warnings': True,
        'progress_hooks': [hook],
        # YouTube 現在要靠 JavaScript 執行環境解驗證；預設只找 deno，這台電腦有 Node.js，一起開
        'js_runtimes': {'deno': {}, 'node': {}},
    }
    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(url, download=True)
        path = ydl.prepare_filename(info)
    return path, info


# ── 標點與斷句（v1.8，A1–A3，不用 LLM）──────────────────────────────────────
# A2：Whisper 比較會「模仿範例」而不是「聽說明」，所以提示詞給一小段標點正確的範例句。
#     只給中文；英文 Whisper 本來就會下標點，給了反而可能干擾。
ZH_PROMPT = '嗯，大家好。今天我們來聊一聊，好嗎？其實沒有那麼難！我覺得：先試試看，再說。'
CJK_RE = re.compile(r'[㐀-鿿豈-﫿]')
CJK = r'㐀-鿿豈-﫿'
SENT_END = '。！？!?…'
SOFT_END = '，、,；;：:'
# A3：斷句上限。字幕一行讀得完的長度；中文以字數算，英文以字元算
MAX_CHARS = {'cjk': 22, 'other': 84}
MAX_DUR = 7.0      # 秒
PAUSE = 0.6        # 停頓超過這麼久就斷句
LONG_PAUSE = 1.0   # 中文段尾沒標點時：停這麼久補「。」，否則補「，」


def fix_punct(text):
    """A1：中文標點整理。半形改全形、中文字之間的空白改「，」、清掉重複標點。
    不含中文的句子只整理空白（英文 Whisper 本來就處理得好）。"""
    t = re.sub(r'\s+', ' ', text).strip()
    if not CJK_RE.search(t):
        return t
    t = t.replace('...', '…').replace('。。。', '…')
    t = re.sub(r'(?<!\d),|,(?!\d)', '，', t)              # 1,000 這種數字不動
    t = re.sub(r'(?<!\d):|:(?!\d)', '：', t)               # 12:30 不動
    t = t.replace('?', '？').replace('!', '！').replace(';', '；')
    t = re.sub(rf'(?<=[{CJK}])\.(?!\d)', '。', t)           # 中文後面的句點
    t = re.sub(rf'(?<=[{CJK}])\s+(?=[{CJK}])', '，', t)    # Breeze 用空白斷句
    t = re.sub(r'\s*([，。！？、；：…])\s*', r'\1', t)       # 全形標點旁不留空白
    t = re.sub(r'([，、；：])[，、；：]+', r'\1', t)          # ，， → ，
    t = re.sub(r'[，、；：]+([。！？…])', r'\1', t)           # ，。 → 。
    return t.lstrip('，。、；：')


def _split_long(words, max_chars, cjk):
    """太長的一段：往回找最後一個逗號處切；找不到就在中間切。"""
    text_len = lambda ws: len(''.join(w['w'] for w in ws).strip()) if cjk else len(' '.join(w['w'].strip() for w in ws))
    for i in range(len(words) - 1, 0, -1):
        if words[i - 1]['w'].strip()[-1:] in SOFT_END and text_len(words[:i]) >= max_chars * 0.4:
            return words[:i], words[i:]
    mid = max(1, len(words) // 2)
    return words[:mid], words[mid:]


def resegment(raw_segments, cjk):
    """A3：用逐字時間戳把「太長的段」切開。raw_segments 是 faster-whisper 的 Segment（含 words）。
    只在每個原始段落「裡面」切，不跨段合併：Whisper 自己的分段通常已經是句子
    （英文尤其好），真正的問題是 Breeze 那種 30 秒一大段。
    段內遇句尾標點、停頓 ≥ PAUSE 秒就切；超過字數／秒數上限時優先在逗號處切。"""
    max_chars = MAX_CHARS['cjk' if cjk else 'other']
    sent_end = SENT_END if cjk else SENT_END + '.'
    joined = lambda ws: ''.join(w['w'] for w in ws) if cjk else ''.join(w['w'] for w in ws).strip()
    segs = []
    for s in raw_segments:
        ws = []
        for w in (s.words or []):
            txt = w.word
            # Breeze／中文：詞前面的空白代表停頓斷句，標在前一個字後面當逗號
            if cjk and ws and txt[:1].isspace() and CJK_RE.match(txt.strip()[:1] or 'a') \
                    and CJK_RE.search(ws[-1]['w'][-1:]):
                ws[-1]['w'] += '，'
            ws.append({'w': txt.strip() if cjk else txt, 's': w.start, 'e': w.end})
        if ws:
            segs.append(ws)
    if not segs:
        return None
    merged = []                             # [(words, 跟下一段的間隔秒數)]
    for si, words in enumerate(segs):
        next_start = segs[si + 1][0]['s'] if si + 1 < len(segs) else None
        groups, cur = [], []
        for i, w in enumerate(words):
            cur.append(w)
            nxt = words[i + 1] if i + 1 < len(words) else None
            if nxt is None:
                groups.append(cur); break
            gap = nxt['s'] - w['e']
            if w['w'].strip()[-1:] in sent_end or gap >= PAUSE:
                groups.append(cur); cur = []
            elif len(joined(cur)) >= max_chars or cur[-1]['e'] - cur[0]['s'] >= MAX_DUR:
                head, tail = _split_long(cur, max_chars, cjk)
                groups.append(head); cur = list(tail)
        # 段內太短的（一兩個字、不到 0.8 秒）併回前一塊
        pieces = []
        for g in groups:
            t = joined(g)
            tiny = (len(t) <= 2 if cjk else len(t.split()) <= 1) or g[-1]['e'] - g[0]['s'] < 0.8
            if pieces and tiny:
                pieces[-1] = pieces[-1] + g
            else:
                pieces.append(g)
        for pi, g in enumerate(pieces):
            if pi + 1 < len(pieces):
                gap = pieces[pi + 1][0]['s'] - g[-1]['e']
            else:
                gap = (next_start - g[-1]['e']) if next_start is not None else 99
            # 跨段：Whisper 自己把一句話尾巴切成很小一段（"…but they don't" / "know."）時，
            # 前一行還沒講完（沒有句尾標點）又沒停頓，就把小尾巴接回去
            t = joined(g)
            tiny = (len(t) <= 2 if cjk else len(t.split()) <= 1)
            if pi == 0 and tiny and merged and merged[-1][1] < PAUSE \
                    and merged[-1][0][-1]['w'].strip()[-1:] not in sent_end:
                merged[-1] = (merged[-1][0] + g, gap)
            else:
                merged.append((g, gap))
    out = []
    for ws, gap in merged:
        text = fix_punct(joined(ws))
        if not text:
            continue
        if cjk and text[-1] not in SENT_END + '，、；：':
            text += '。' if gap >= LONG_PAUSE else '，'
        if cjk and text.endswith('，') and gap >= LONG_PAUSE:
            text = text[:-1] + '。'
        out.append({'start': round(ws[0]['s'], 2), 'end': round(ws[-1]['e'], 2), 'text': text})
    return out


def split_by_text(start, end, text, cjk):
    """沒有逐字時間戳時的斷句（台語 Breeze 用這個）：先整理標點（Breeze 的空白會變成「，」），
    再照標點切成子句、依字數上限打包；每塊的時間照字數比例分配（大約準，夠用）。"""
    text = fix_punct(text)
    max_chars = MAX_CHARS['cjk' if cjk else 'other']
    ends = SENT_END + ('' if cjk else '.')
    clauses = [c for c in re.split(r'(?<=[' + re.escape(ends + SOFT_END) + r'])\s*', text) if c]
    pieces, cur = [], ''
    for c in clauses:
        if cur and len(cur) + len(c) > max_chars:
            pieces.append(cur); cur = ''
        cur += c if cjk else (' ' + c if cur else c)
        if cur[-1:] in ends:
            pieces.append(cur); cur = ''
    if cur:
        pieces.append(cur)
    total = sum(len(p) for p in pieces) or 1
    out, t = [], start
    for i, p in enumerate(pieces):
        t2 = end if i == len(pieces) - 1 else t + (end - start) * len(p) / total
        if cjk and p[-1] in '，、；：' and i == len(pieces) - 1:
            p = p[:-1] + '。'
        elif cjk and p[-1] not in SENT_END + '，、；：':
            p += '。' if i == len(pieces) - 1 else '，'
        out.append({'start': round(t, 2), 'end': round(t2, 2), 'text': p})
        t = t2
    return out


def transcribe_file(model, path, language, output, taigi, on_segment=None):
    """辨識一個音檔，回傳 (segments, 偵測到的語言, 長度秒數)。
    language: 'auto' 或語言代碼；output: 'orig' | 'en' | 'both'（這裡只處理 en 的 task）。"""
    kwargs = dict(beam_size=5, vad_filter=True)
    if taigi:
        kwargs['language'] = 'zh'           # Breeze 用 Whisper 的 zh token 輸出華語
    elif language != 'auto':
        kwargs['language'] = language
    else:
        # 自動偵測時先偵測語言，才知道要不要給中文範例提示
        from faster_whisper import decode_audio
        lang, _, _ = model.detect_language(decode_audio(path, sampling_rate=16000))
        kwargs['language'] = lang
    if output == 'en':
        kwargs['task'] = 'translate'        # Whisper 內建只能翻成英文
    elif kwargs['language'] == 'zh' and not taigi:
        # Breeze 不給提示：實測給了範例提示，時間戳會壞掉（6 段都變成 1.5–1.6 秒）、用字也變
        kwargs['initial_prompt'] = ZH_PROMPT
    translating = kwargs.get('task') == 'translate'
    cjk = kwargs['language'] in ('zh', 'ja', 'yue') and not translating
    # A3 斷句用逐字時間戳。台語 Breeze 不開：實測開了之後慢很多倍，結尾還會在配樂上
    # 「編」出不存在的句子（「你就知道有多好笑了吧」），改用 split_by_text 依標點切
    use_words = not taigi
    if use_words:
        kwargs['word_timestamps'] = True
        # 跳過超過 2 秒沒有人聲時冒出的內容（逐字時間戳模式才有作用）
        kwargs['hallucination_silence_threshold'] = 2.0
    segments, info = model.transcribe(path, **kwargs)
    raw = []
    for s in segments:                      # 逐段收集，讓進度條動
        raw.append(s)
        if on_segment:
            on_segment(s, info.duration)
    out = resegment(raw, cjk) if use_words else None
    if out is None:                         # 沒有逐字時間戳：依標點切開、時間照字數比例
        out = []
        for i, s in enumerate(raw):
            end = s.end
            # 保險：模型偶爾給出「很長的字、極短的時間」（每字不到 0.05 秒），
            # 就把結束時間延到下一段開始（最後一段延到音檔結尾）
            if end - s.start < 0.05 * len(s.text.strip()):
                end = raw[i + 1].start if i + 1 < len(raw) else (info.duration or end)
            out.extend(split_by_text(s.start, end, s.text, cjk))
    return out, info.language, info.duration


def translate_segments(model, path, out, src_lang, job):
    """並列模式第二階段：把每段原文對應的那一小段聲音丟去翻成英文，
    英文自然跟原文一句對一句（分開跑兩次辨識的話，兩邊的分段長短對不起來）。"""
    from faster_whisper import decode_audio
    audio = decode_audio(path, sampling_rate=16000)
    for i, seg in enumerate(out):
        a, b = int(seg['start'] * 16000), int(seg['end'] * 16000)
        clip = audio[a:b]
        if len(clip) < 16000 * 0.3:   # 太短的片段翻不出東西，還容易冒出 "Thank you."
            seg['en'] = ''
        else:
            parts, _ = model.transcribe(clip, task='translate', language=src_lang, beam_size=5,
                                        vad_filter=False, without_timestamps=True,
                                        condition_on_previous_text=False)
            seg['en'] = ' '.join(p.text.strip() for p in parts).strip()
        job['progress'] = 50 + int((i + 1) * 49 / len(out))


def run_job(job, url, model_name, language, output='orig'):
    path = None
    try:
        job.update(status='downloading', message='下載影片音軌中…', progress=0)
        path, info = download_audio(url, job)
        # FB 的 title 前面會黏「5.3K views · 109 reactions | 」之類的計數（可能好幾個），全部去掉
        title = info.get('title') or info.get('description') or ''
        job['title'] = re.sub(r'^(?:[\d.,KkMm萬]+\s*(?:reactions?|views?|個心情|次觀看)\s*[|·]\s*)+', '', title)
        job['duration'] = info.get('duration') or 0

        taigi = language == 'nan'
        if taigi:
            model_name = TAIGI_MODEL_DIR
            job['model'] = 'Breeze-ASR-26 (台語→華語)'
            if not os.path.isfile(os.path.join(TAIGI_MODEL_DIR, 'model.bin')):
                raise RuntimeError('找不到台語模型，請先照 README「台語模型」下載 Breeze-ASR-26 int8')
        job.update(status='loading', message=f'載入模型 {job["model"]}（第一次使用需下載）…', progress=0)
        model = get_model(model_name)

        job.update(status='transcribing', message='語音辨識中…', progress=0)
        if taigi:
            output = 'orig'           # Breeze 不支援翻譯
        # 並列模式：辨識佔進度 0–50%，翻譯佔 50–100%
        span = 50 if output == 'both' else 99
        preview = []

        def on_segment(s, total):
            # 辨識中先顯示原始分段（標點已整理），完成後換成重新斷句的版本
            preview.append({'start': round(s.start, 2), 'end': round(s.end, 2), 'text': fix_punct(s.text)})
            job['segments'] = preview
            job['progress'] = min(span, int(s.end * span / (total or job['duration'] or 1)))

        out, lang, _ = transcribe_file(model, path, language, output, taigi, on_segment)
        job['segments'] = out
        job['language'] = '台語（輸出華語）' if taigi else lang
        if output == 'en':
            job['language'] += ' → 英文'
        both = output == 'both'
        if both and lang == 'en':
            both = False              # 原文就是英文，不用再翻
            job['note'] = '原文已經是英文，所以沒有另外翻譯'
        if both and out:
            job.update(message='翻譯成英文中…', progress=50)
            job['language'] += ' ＋ 英文'
            translate_segments(model, path, out, lang, job)
        job.update(status='done', message='完成', progress=100)
    except Exception as e:
        job.update(status='error', message=f'失敗：{e}', error=str(e))
    finally:
        if path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
        job['elapsed'] = round(time.time() - job['started'], 1)


@app.get('/')
def index():
    return send_from_directory(HERE, 'index.html')


@app.post('/api/transcribe')
def transcribe():
    data = request.get_json(silent=True) or {}
    url = (data.get('url') or '').strip()
    model_name = data.get('model') or 'small'
    # 預設自動偵測：強制 zh 時，英文影片會被 Whisper 直接「翻譯」成中文
    language = data.get('language') or 'auto'
    if not URL_RE.match(url):
        return jsonify(error='請輸入 Facebook 或 YouTube 影片網址'), 400
    if model_name not in MODELS:
        return jsonify(error='未知的模型'), 400
    if language not in ('auto', 'zh', 'en', 'ja', 'fr', 'nan'):
        return jsonify(error='未知的語言'), 400
    # output: orig 只要原文 / en 只要英文 / both 原文＋英文並列（舊版的 translate:true 視同 en）
    output = data.get('output') or ('en' if data.get('translate') else 'orig')
    if output not in ('orig', 'en', 'both'):
        return jsonify(error='未知的輸出方式'), 400
    job_id = uuid.uuid4().hex[:12]
    job = {'id': job_id, 'status': 'queued', 'message': '排隊中…', 'progress': 0,
           'segments': [], 'started': time.time(), 'url': url, 'model': model_name}
    jobs[job_id] = job
    threading.Thread(target=run_job, args=(job, url, model_name, language, output), daemon=True).start()
    return jsonify(id=job_id)


@app.post('/api/docx')
def export_docx():
    """把逐字稿輸出成 Word (.docx)。內容由前端送來（使用者可能已在頁面上看過）。"""
    import io
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from flask import send_file

    data = request.get_json(silent=True) or {}
    segments = data.get('segments') or []
    show_ts = bool(data.get('timestamps'))

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(12)
    # 中文字型要另外設 eastAsia，不然 Word 會用預設的新細明體以外字型亂套
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')

    doc.add_heading(data.get('title') or '影片逐字稿', level=1)
    for line in (data.get('url'), data.get('meta')):
        if line:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x65, 0x67, 0x6B)

    for s in segments:
        p = doc.add_paragraph()
        if show_ts:
            t = float(s.get('start') or 0)
            ts = p.add_run(f'[{int(t // 60):02d}:{int(t % 60):02d}] ')
            ts.font.color.rgb = RGBColor(0x8A, 0x8D, 0x91)
        p.add_run(str(s.get('text') or ''))
        if s.get('en'):
            p.add_run().add_break()
            en = p.add_run(str(s['en']))
            en.italic = True
            en.font.size = Pt(10.5)
            en.font.color.rgb = RGBColor(0x65, 0x67, 0x6B)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    name = re.sub(r'[^\w.-]', '_', data.get('filename') or 'transcript') + '.docx'
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@app.get('/api/job/<job_id>')
def job_status(job_id):
    job = jobs.get(job_id)
    if not job:
        return jsonify(error='找不到工作'), 404
    return jsonify({k: v for k, v in job.items() if k != 'started'})


def main():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 8792
    print(f'FB／YouTube 影片轉文字 -> http://localhost:{port}/   （Ctrl+C 結束）')
    app.run(host='127.0.0.1', port=port, threaded=True)


if __name__ == '__main__':
    main()
